#!/usr/bin/env python3
"""
Materialize xlsx fixtures used by tests/ingest_xlsx_test.bats and
tests/watchdog_xlsx_test.bats. Run manually after editing — fixtures
are committed binaries; CI does not regenerate them.

Requires: pip3 install openpyxl
"""
from __future__ import annotations
from pathlib import Path
from openpyxl import Workbook

FIXTURES = Path(__file__).resolve().parent


def write_single_sheet() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["region", "amount", "date"])
    ws.append(["EMEA", 100, "2026-01-01"])
    ws.append(["AMER", 250, "2026-01-02"])
    ws.append(["APAC", 75, "2026-01-03"])
    wb.save(FIXTURES / "single-sheet.xlsx")


def write_multi_with_hidden() -> None:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Sales"
    ws1.append(["region", "amount"])
    ws1.append(["EMEA", 100])
    ws2 = wb.create_sheet("Inventory")
    ws2.append(["sku", "qty"])
    ws2.append(["A1", 12])
    ws2.append(["B2", 7])
    ws3 = wb.create_sheet("HiddenScratch")
    ws3.append(["x", "y"])
    ws3.append([1, 2])
    ws3.sheet_state = "hidden"
    wb.save(FIXTURES / "multi-sheet-with-hidden.xlsx")


def write_empty_and_hidden() -> None:
    wb = Workbook()
    wb.active.title = "Empty"
    extra = wb.create_sheet("AlsoHidden")
    extra.append(["only", "header"])
    extra.sheet_state = "hidden"
    wb.save(FIXTURES / "empty-and-hidden.xlsx")


def write_collision() -> None:
    # Two sheets that kebab-slugify to the same string within the workbook.
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Sales Data"
    ws1.append(["region", "amount"])
    ws1.append(["EMEA", 100])
    ws2 = wb.create_sheet("sales-data")
    ws2.append(["region", "amount"])
    ws2.append(["AMER", 250])
    wb.save(FIXTURES / "collision-sheet-names.xlsx")


def write_corrupt() -> None:
    # Truncated zip header — calamine fails to parse.
    (FIXTURES / "corrupt.xlsx").write_bytes(b"PK\x03\x04not-a-real-xlsx")


def main() -> None:
    write_single_sheet()
    write_multi_with_hidden()
    write_empty_and_hidden()
    write_collision()
    write_corrupt()
    print("OK|fixtures regenerated under", FIXTURES)


if __name__ == "__main__":
    main()
